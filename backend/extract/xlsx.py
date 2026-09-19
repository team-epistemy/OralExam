"""XLSX extractor: one unit per worksheet, rows rendered as tab-separated text.

openpyxl is pure-Python (no system libs). Each sheet becomes a structural unit
whose heading_path is the sheet name, so the chunker keeps sheets separate.
"""
from __future__ import annotations
import io
from typing import List

from backend.models import ChunkPosition
from backend.chunking import ExtractedUnit


def _rows_to_text(rows) -> str:
    """Render rows as TSV, dropping fully empty rows."""
    lines = []
    for row in rows:
        cells = ["" if v is None else str(v) for v in row]
        if any(c.strip() for c in cells):
            lines.append("\t".join(cells).rstrip("\t"))
    return "\n".join(lines)


class XlsxExtractor:
    """openpyxl read-only walk; one ExtractedUnit per non-empty worksheet."""

    def extract(self, data: bytes) -> List[ExtractedUnit]:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        units: List[ExtractedUnit] = []
        for ws in wb.worksheets:
            text = _rows_to_text(ws.iter_rows(values_only=True))
            if text.strip():
                units.append(ExtractedUnit(
                    text=text, position=ChunkPosition(heading_path=[ws.title])))
        wb.close()
        return units
